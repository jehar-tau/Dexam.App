"""Export the canonical Markdown role guides to a printable Word handbook.

Run with the bundled document runtime (python-docx). Render using the documents
skill's render_docx.py --emit_pdf, inspect every page, then copy the verified PDF
alongside the Word file. No network, app, database or provider access is used.
"""

import re
from pathlib import Path
from urllib.parse import urljoin

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
GUIDES = ROOT / "docs/usage"
OUT = GUIDES / "downloads"
BASE = "https://github.com/jehar-tau/Dexam.App/blob/main/docs/usage/"
FILES = ["README.md", "ACCESS_AND_TROUBLESHOOTING.md", "STUDENTS.md", "TEACHERS.md", "STAFF.md", "ADMINS.md"]
LABELS = ["Using this handbook", "Access and troubleshooting", "Student guide", "Teacher guide", "Staff guide", "Administrator guide"]


def plain(value):
    return value.replace("—", " - ").replace("–", "-").replace("→", " to ").replace("…", "...")


def link(p, label, target):
    el = OxmlElement("w:hyperlink")
    rel = p.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    el.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "244B70")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = plain(label)
    run.append(text)
    el.append(run)
    p._p.append(el)


def inline(p, text):
    for token in re.split(r"(\[[^]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)", text):
        match = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token)
        if match:
            link(p, match[1], urljoin(BASE, match[2]))
        elif token.startswith("**") and token.endswith("**"):
            p.add_run(plain(token[2:-2])).bold = True
        elif token.startswith("`") and token.endswith("`"):
            p.add_run(plain(token[1:-1]))
        else:
            p.add_run(plain(token))


doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.78)
sec.footer_distance = Inches(.3)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.03
doc.styles["Title"].font.size = Pt(30)
doc.styles["Subtitle"].font.size = Pt(14)
for name, size in [("Heading 1", 22), ("Heading 2", 15), ("Heading 3", 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.keep_with_next = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Dexam application handbook  |  ").font.size = Pt(9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Dexam application handbook"
doc.core_properties.author = "Dexam"
doc.core_properties.subject = "Application usage for students teachers staff and administrators"

doc.add_paragraph("Dexam application handbook", "Title")
doc.add_paragraph("Students teachers staff and administrators", "Subtitle")
doc.add_paragraph("Edition 18 October 2026")
doc.add_paragraph("Use this handbook to sign in, find your permitted workspace and complete the application's existing workflows. Read the access guide first, then the section for your responsibilities. The guides distinguish available functionality from services that remain disabled or require launch approval.")
doc.add_paragraph("This edition describes the fictional pilot. It does not authorize real student onboarding or production operations. Portal selection does not grant permissions.")
p = doc.add_paragraph()
link(p, "Read the current handbook online", "https://github.com/jehar-tau/Dexam.App/tree/main/docs/usage")
doc.add_heading("Contents", 2)
for i, label in enumerate(LABELS, 1):
    doc.add_paragraph(f"{i}  {label}")
doc.add_paragraph("Printed copies are dated snapshots. Check the online handbook before relying on a route, limit, permission or service availability. Keep credentials and private student information out of this document.")

for index, filename in enumerate(FILES):
    if index == 0:
        doc.add_page_break()
    doc.add_heading(f"{index + 1}  {LABELS[index]}", 1)
    lines = (GUIDES / filename).read_text().splitlines()
    pos = 0
    while pos < len(lines):
        line = lines[pos].strip()
        pos += 1
        if not line or line.startswith("# ") or line.startswith("[Handbook home]"):
            continue
        # Download/maintenance metadata is online-only and must not recursively
        # include the generated handbook in its own export.
        if line == "## Printable and online editions":
            while pos < len(lines) and not lines[pos].startswith("## "):
                pos += 1
            continue
        if line.startswith("|"):
            rows = [line]
            while pos < len(lines) and lines[pos].strip().startswith("|"):
                rows.append(lines[pos].strip())
                pos += 1
            headers = [v.strip() for v in rows[0].strip("|").split("|")]
            # Preserve every cell as labelled records, avoiding wide tables in print.
            for row in rows[2:]:
                cells = [v.strip() for v in row.strip("|").split("|")]
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                inline(p, "**" + cells[0] + "**")
                for header, cell in zip(headers[1:], cells[1:]):
                    p = doc.add_paragraph()
                    inline(p, (header + ": " if len(headers) > 2 else "") + cell)
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)", line)
        if heading:
            title = re.sub(r"[^\w\s]", " ", plain(heading[2]))
            doc.add_heading(re.sub(r"\s+", " ", title).strip(), len(heading[1]))
            continue
        bullet = re.match(r"^-\s+(.+)", line)
        numbered = re.match(r"^(\d+)\.\s+(.+)", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            inline(p, bullet[1])
        elif numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.22)
            p.paragraph_format.first_line_indent = Inches(-.22)
            inline(p, numbered[1] + ". " + numbered[2])
        else:
            paragraph = line
            while pos < len(lines) and lines[pos].strip() and not re.match(r"^(#|\||- |\d+\. )", lines[pos]):
                paragraph += " " + lines[pos].strip()
                pos += 1
            p = doc.add_paragraph()
            inline(p, paragraph)

OUT.mkdir(parents=True, exist_ok=True)
target = OUT / "Dexam-application-handbook.docx"
doc.save(target)
print(target)
